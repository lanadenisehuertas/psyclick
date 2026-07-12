from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from PIL import Image, ImageDraw, ImageFont


SOURCE = Path(r"C:\Users\Lana\Downloads\PsyClick_CS0029_Security_Document.docx")
ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "artifacts" / "PsyClick_CS0029_Security_Document_COMPLETED.docx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17324D"
TEAL = "0A8F8C"
LIGHT_TEAL = "DDEFEF"
LIGHT_BLUE = "EAF1F7"
LIGHT_GRAY = "F2F4F6"
WHITE = "FFFFFF"
RED = "B42318"
AMBER = "9A6700"
GREEN = "067647"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_table(table, header_fill=NAVY, widths=None):
    table.style = "Table Grid"
    table.autofit = True
    for i, cell in enumerate(table.rows[0].cells):
        shade(cell, header_fill)
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.name = "Arial"
                run.font.size = Pt(8.5)
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
    for row_index, row in enumerate(table.rows[1:], start=1):
        if row_index % 2 == 0:
            for cell in row.cells:
                shade(cell, LIGHT_GRAY)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)


def add_status_color(cell, status):
    normalized = status.lower()
    color = GREEN if "implemented" in normalized and "not" not in normalized else AMBER
    if "not implemented" in normalized or "critical" in normalized or "failed" in normalized:
        color = RED
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph_after(document, anchor, text="", style=None):
    paragraph = document.add_paragraph(text, style=style)
    anchor._p.addnext(paragraph._p)
    return paragraph


def add_table_after(document, anchor, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color=WHITE)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            set_cell_text(cells[idx], value)
    format_table(table, widths=widths)
    anchor._p.addnext(table._tbl)
    return table


def replace_table_rows(table, headers, rows):
    while len(table.columns) < len(headers):
        table.add_column(Inches(1.4))
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color=WHITE)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            set_cell_text(cells[idx], value)
    format_table(table)


def replace_paragraph_starting(document, prefix, text):
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(prefix):
            paragraph.text = text
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def replace_paragraph_after_heading(document, heading, text):
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() == heading:
            paragraphs[index + 1].text = text
            return paragraphs[index + 1]
    raise ValueError(f"Heading not found: {heading}")


def find_table(document, first_header, second_header=None):
    for table in document.tables:
        if not table.rows or not table.rows[0].cells:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers[0] == first_header and (second_header is None or (len(headers) > 1 and headers[1] == second_header)):
            return table
    raise ValueError(f"Table not found: {first_header} / {second_header}")


def insert_paragraph_sequence_after(document, anchor, items):
    """Insert [(text, style), ...] in order immediately after anchor."""
    previous = anchor
    inserted = []
    for text, style in items:
        paragraph = document.add_paragraph(text, style=style)
        previous._p.addnext(paragraph._p)
        previous = paragraph
        inserted.append(paragraph)
    return inserted


def create_diagrams(directory):
    directory.mkdir(parents=True, exist_ok=True)
    paths = {}
    try:
        font = ImageFont.truetype("arial.ttf", 22)
        small = ImageFont.truetype("arial.ttf", 14)
    except OSError:
        font = ImageFont.load_default(); small = ImageFont.load_default()

    def make(name, title, boxes, links, width=1500, height=760):
        image = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(image)
        draw.text((width // 2, 25), title, fill="#17324D", font=font, anchor="ma")
        for x, y, w, h, label, fill in boxes:
            draw.rounded_rectangle((x, y, x+w, y+h), radius=16, fill=fill, outline="#17324D", width=3)
            draw.multiline_text((x+w//2, y+h//2), label, fill="#17324D", font=small, anchor="mm", align="center")
        for x1, y1, x2, y2, label in links:
            draw.line((x1, y1, x2, y2), fill="#0A8F8C", width=3)
            draw.polygon([(x2, y2), (x2-12, y2-7), (x2-12, y2+7)], fill="#0A8F8C")
            if label:
                draw.text(((x1+x2)//2, (y1+y2)//2-12), label, fill="#17324D", font=small, anchor="mm")
        path = directory / name
        image.save(path)
        return path

    paths["architecture"] = make("system-architecture.png", "PsyClick Layered Security Architecture", [
        (60,120,260,90,"Clinician / Patient UI","#DDEFEF"),(370,120,260,90,"Auth + Consent\nSession Manager","#EAF1F7"),(680,120,260,90,"Telemetry + HAL\nFeature Extraction","#EAF1F7"),(990,120,260,90,"Analytics Engine\nEWMA → T² → PSI/PAI","#EAF1F7"),
        (260,350,260,90,"Report Generator\nHTML / PDF","#F5EEDC"),(620,350,260,90,"SQLite Datastore\nclinical + audit","#F5EEDC"),(980,350,260,90,"Protected Backup\n(encrypted target)","#F5EEDC"),
        (80,590,260,90,"Security Boundary\nRBAC • encryption • audit","#F7E5E5"),(460,590,260,90,"Endpoint Boundary\nOS ACL • screen lock","#F7E5E5"),(840,590,260,90,"Availability Boundary\nbackup • restore • monitor","#F7E5E5")], [
        (320,165,370,165,"events / consent"),(630,165,680,165,"authorized capture"),(940,165,990,165,"features"),(1090,210,750,350,"results"),(750,395,520,395,"report data"),(750,440,750,350,"ACID records"),(880,395,980,395,"verified backup"),(210,590,390,440,""),(590,590,700,440,""),(970,590,1110,440,"")])
    paths["dfd"] = make("dfd.png", "PsyClick Context and Level-1 Data Flow", [
        (30,150,210,90,"Clinician","#DDEFEF"),(30,520,210,90,"Patient /\nParticipant","#DDEFEF"),(1250,150,210,90,"Administrator","#DDEFEF"),(1250,520,210,90,"Auditor / DPO","#DDEFEF"),
        (360,140,240,95,"P1 Auth +\nAccess Control","#EAF1F7"),(360,300,240,95,"P2 Intake +\nConsent","#EAF1F7"),(360,460,240,95,"P3 Capture + HAL","#EAF1F7"),(760,300,270,95,"P4/P5 Features +\nT² Classification","#EAF1F7"),(760,460,270,95,"P6 Report +\nP7 Security Ops","#F5EEDC")], [
        (240,195,360,185,"credentials / requests"),(240,565,360,505,"events / consent"),(600,185,760,345,"authorized context"),(600,345,760,505,"normalized data"),(1030,345,1250,195,"status / backup"),(1030,505,1250,565,"audit reports")])
    paths["erd"] = make("erd.png", "Core Entity-Relationship Model", [
        (35,130,260,105,"users\nPK user_id\nusername • role_id","#EAF1F7"),(35,470,260,105,"roles\nPK role_id\nrole_name","#EAF1F7"),(370,130,260,105,"patients\nPK patient_id\npseudonym","#EAF1F7"),(370,470,260,105,"consents\nPK consent_id\ndecision • timestamp","#EAF1F7"),(705,130,260,105,"sessions\nPK session_id\npatient • clinician","#EAF1F7"),(705,470,260,105,"analysis_results\nPK result_id\nT² • PSI • PAI • flag","#EAF1F7"),(1040,130,260,105,"reports\nPK report_id\nfile_hash","#EAF1F7"),(1040,470,260,105,"audit_logs\nPK audit_id\nentry_hash","#EAF1F7")], [
        (295,180,370,180,"role"),(630,180,705,180,"patient"),(500,235,500,470,"consent"),(835,235,835,470,"result"),(965,180,1040,180,"report"),(965,520,1040,520,"audit")])
    paths["network"] = make("network-architecture.png", "Local-First Network and Trust Boundaries", [
        (50,220,260,100,"Patient Interaction\nkeyboard • mouse • display","#DDEFEF"),(360,220,260,100,"Electron Desktop\nReact UI","#EAF1F7"),(670,220,260,100,"Loopback API\n127.0.0.1:5001","#EAF1F7"),(1030,220,220,100,"SQLite\nlocal data","#F5EEDC"),(250,500,290,90,"Endpoint controls\nOS ACL • lock • EDR","#F7E5E5"),(700,500,310,90,"Protected backup target\nencrypted removable/LAN","#F7E5E5")], [
        (310,270,360,270,"supervised"),(620,270,670,270,"HTTP loopback"),(930,270,1030,270,"SQL"),(800,320,850,500,"verified backup")])
    return paths


document = Document(SOURCE)
document.core_properties.title = "PsyClick CS0029 Information Assurance and Security Document"
document.core_properties.subject = "Completed security architecture, risk assessment, and as-built evaluation"
document.core_properties.author = "Añonuevo; Ballano; Huertas; Tablate"
document.core_properties.comments = (
    "Completed against the CS0029 Final Project Specification, the PsyClick thesis, "
    "and a static review of the system workspace on 10 July 2026."
)

# Add the thesis-required Chapter 1 and Chapter 2 coverage before the security-specific chapters.
overview_heading = next(p for p in document.paragraphs if p.text.strip() == "PROJECT OVERVIEW")
cover_date = next(p for p in document.paragraphs if p.text.strip() == "July 2026")
chapter12 = [
    ("CHAPTER 1 – INTRODUCTION", "Heading 1"),
    ("1.1 Background of the Study", "Heading 2"),
    ("Mental-health assessment in educational and counseling settings often depends on verbal self-disclosure and clinician observation. The thesis identifies concealment, recall bias, self-stigma, and delayed help-seeking as barriers that can cause clinically relevant distress to be missed. PsyClick addresses this gap by collecting consent-gated keyboard and cursor timing behavior during a controlled single-session workflow and converting it into transparent, clinician-facing psychomotor indicators. The system is intended to augment professional assessment, not replace diagnosis.", "Normal"),
    ("1.2 Problem Statement", "Heading 2"),
    ("The project addresses the need for objective, within-session evidence of psychomotor slowing or agitation while protecting sensitive behavioral and mental-health-related information. The security problem is equally important: behavioral telemetry can act as a biometric signal, so unauthorized access, disclosure, tampering, weak authentication, consent bypass, unsafe export, and data loss could harm patients and undermine clinical trust.", "Normal"),
    ("1.3 Objectives", "Heading 2"),
    ("The study develops a hardware-abstracted telemetry module, a within-session EWMA baseline, an eight-feature Hotelling’s T² anomaly engine, PSI/PAI feature contribution analysis, fuzzy classification, and Green/Amber/Red decision-support outputs. The CS0029 security objective is to protect confidentiality, integrity, availability, authentication, authorization, accountability, privacy, and clinical safety across that workflow.", "Normal"),
    ("1.4 Scope and Limitations", "Heading 2"),
    ("PsyClick covers a local desktop workflow consisting of clinician login, pseudonymous intake, consent, keyboard and mouse calibration, PHQ-9/GAD-7 and emotional tasks, analytics, dashboard/report review, and audit-event capture. The thesis evaluated 105 normative adult participants and 15 clinical sessions. The system is non-diagnostic and is not a replacement for clinical judgment. The supplied prototype is not evidence of production-grade password hashing, server-side RBAC, database encryption, tamper-evident logging, or tested backup/recovery; those items are explicitly treated as security hardening requirements in this document.", "Normal"),
    ("1.5 Significance of the Study", "Heading 2"),
    ("Mental-health professionals receive an objective behavioral signal that can complement interviews and self-report instruments. Participants benefit from a structured, consent-based process that does not require semantic analysis of typed responses. Researchers receive a reproducible white-box pipeline for hardware normalization, personal baselines, multivariate detection, and interpretable reporting. Institutions receive a concrete information-assurance baseline for handling behavioral and clinical data locally.", "Normal"),
    ("CHAPTER 2 – REVIEW OF RELATED LITERATURE AND STUDIES", "Heading 1"),
    ("2.1 Local Literature", "Heading 2"),
    ("The thesis situates PsyClick in the Philippine context, where educational and community mental-health services face access constraints, stigma, and limited clinician capacity. Local literature supports the need for earlier, complementary assessment evidence while emphasizing ethical safeguards, privacy, and the continued role of trained professionals.", "Normal"),
    ("2.2 Foreign Literature", "Heading 2"),
    ("Foreign literature reviewed in the thesis supports keystroke dynamics, mouse movement analysis, hardware abstraction, EWMA monitoring, Hotelling’s T², shrinkage covariance estimation, fuzzy classification, and human-centered clinical decision support. Together, these sources justify combining timing, motor, and contextual features rather than relying on a single speed or reaction-time measure.", "Normal"),
    ("2.3 Local Studies", "Heading 2"),
    ("Local studies establish the relevance of digital and technology-assisted mental-health services in Philippine educational settings, while also highlighting adoption, privacy, and usability constraints. PsyClick extends that work by implementing a concrete local-first prototype and evaluating software quality and clinician agreement.", "Normal"),
    ("2.4 Foreign Studies", "Heading 2"),
    ("Foreign studies demonstrate that keystroke and cursor behavior can provide measurable correlates of stress, workload, depression-related psychomotor change, and cognitive load. The thesis synthesizes these findings into a single-session, personalized baseline model and adds contribution analysis and domain-segmented outputs for interpretability.", "Normal"),
    ("2.5 Synthesis", "Heading 2"),
    ("The literature converges on three design requirements: measurements must be hardware-aware, analysis must account for within-person variation, and outputs must be interpretable and ethically bounded. PsyClick’s HAL, EWMA/T² pipeline, PSI/PAI decomposition, and non-diagnostic dashboard implement those requirements in one workflow. Information assurance adds a fourth requirement: the data lifecycle must be consent-gated, least-privilege, auditable, recoverable, and privacy-preserving.", "Normal"),
    ("2.6 Research Gap", "Heading 2"),
    ("Prior work commonly studies isolated keystroke or mouse features, population-level classification, or laboratory tasks. The thesis identifies a gap in single-session, within-person, multivariate analysis that combines hardware normalization, dynamic baselines, transparent contribution analysis, domain profiling, and clinician-facing safeguards. The CS0029 security work addresses a parallel gap by translating that research prototype into a documented, testable security baseline rather than assuming that functional validity alone establishes information assurance.", "Normal"),
]
insert_paragraph_sequence_after(document, cover_date, chapter12)

# Create and embed the architecture, DFD, ERD, and network diagrams required by the specification.
diagram_paths = create_diagrams(OUTPUT.parent / "diagrams")


def insert_picture_after_heading(document, heading, picture_path, caption):
    anchor = next(p for p in document.paragraphs if p.text.strip() == heading)
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = picture_paragraph.add_run()
    run.add_picture(str(picture_path), width=Inches(6.5))
    anchor._p.addnext(picture_paragraph._p)
    caption_paragraph = document.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph._p.addnext(caption_paragraph._p)


insert_picture_after_heading(document, "3.4.1 Context Diagram – Level 0", diagram_paths["dfd"], "Figure 1. PsyClick context and Level-1 data flow.")
insert_picture_after_heading(document, "3.5 Entity Relationship Diagram (ERD)", diagram_paths["erd"], "Figure 2. Core PsyClick entity-relationship model.")
insert_picture_after_heading(document, "3.6 Network Architecture", diagram_paths["network"], "Figure 3. Local-first network and trust boundaries.")
insert_picture_after_heading(document, "3.7 Security Architecture", diagram_paths["architecture"], "Figure 4. Layered PsyClick security architecture (target hardened baseline).")

# Make the scope and evidence boundary explicit.
document.paragraphs[23].text = (
    "From an Information Assurance and Security perspective, PsyClick processes highly sensitive "
    "behavioral and mental-health-related data. Keystroke and cursor dynamics may function as "
    "behavioral biometrics; therefore, confidentiality, integrity, availability, authentication, "
    "authorization, accountability, privacy, and clinical safety must be addressed throughout the "
    "system life cycle. This document defines the required hardened baseline and separately records "
    "the controls evidenced in the current application."
)
document.paragraphs[24].text = (
    "Document basis and scope: thesis-validated analytical and usability claims are derived from the "
    "BYTEME thesis paper; mandatory deliverables are derived from the CS0029 Final Project "
    "Specification; and implementation status is based on a static review of the supplied PsyClick "
    "workspace and packaged modules on 10 July 2026. No live penetration test or production-data test "
    "was performed. A control is described as implemented only where the reviewed system provides "
    "direct evidence; all other controls are labeled Partial, Not Implemented, Proposed, or Not Tested."
)

# Insert an early, submission-friendly implementation summary.
anchor = document.paragraphs[24]
summary_heading = add_paragraph_after(document, anchor, "AS-BUILT SECURITY IMPLEMENTATION SUMMARY", "Heading 1")
summary_heading.paragraph_format.page_break_before = True
summary_note = add_paragraph_after(
    document,
    summary_heading,
    "The application is a working clinical decision-support prototype, but it does not yet satisfy the "
    "mandatory CS0029 security baseline for institutional deployment. The highest-priority gaps are "
    "plaintext clinician passwords, absence of authenticated server sessions and server-side RBAC, "
    "a UI-only consent gate, unencrypted data and reports, mutable audit records, and the absence of a "
    "verified backup/restore implementation. The target controls in later sections are the remediation "
    "baseline, not claims about the current build.",
)
implementation_rows = [
    ("Authentication", "Partial / Non-compliant", "Clinician login and re-authentication screens exist, but passwords are stored and compared as plaintext; no lockout, MFA, rate limiting, or server session token is present.", "Critical"),
    ("Authorization", "Not Implemented", "The React guard is client-side only. API routes do not enforce an authenticated identity, role, permission, or record ownership.", "Critical"),
    ("Consent", "Partial", "The intake UI disables Start until a checkbox is selected, but the backend start endpoint does not receive or independently validate a persisted consent record.", "High"),
    ("Data protection", "Partial / Non-compliant", "Processing is local and the API binds to loopback, but the SQLite database and HTML exports are not application-encrypted or integrity-hashed.", "High"),
    ("Audit and accountability", "Partial", "Audit events and an audit screen exist, but records are mutable, not hash-chained, and audit read/write endpoints are not authorization-protected.", "High"),
    ("Availability", "Not Implemented", "SQLite transactions provide basic consistency, but no scheduled encrypted backup, restore verification, redundancy, or recovery evidence is implemented.", "High"),
    ("Security monitoring", "Partial", "Failed or unusual events can be written to the audit log, but there is no alerting, rate monitoring, integrity monitor, or formal review workflow.", "Moderate"),
    ("Endpoint isolation", "Partially Implemented", "The Flask API binds to 127.0.0.1 with debug disabled; Electron disables Node integration and enables context isolation. CORS remains unrestricted.", "Moderate"),
]
summary_table = add_table_after(
    document,
    summary_note,
    ["Security Domain", "As-Built Status", "Evidence-Based Finding", "Priority"],
    implementation_rows,
    widths=[1.15, 1.25, 4.0, 0.8],
)
for row in summary_table.rows[1:]:
    add_status_color(row.cells[1], row.cells[1].text)
    add_status_color(row.cells[3], row.cells[3].text)

# Correct claims in the implementation and evaluation narrative.
replace_paragraph_starting(document, "Passwords are never encrypted reversibly.", (
    "Current implementation status: the reviewed database schema stores clinician passwords in a "
    "plaintext password field and the verification routine performs direct string comparison. This is "
    "a Critical finding and does not meet secure password-storage requirements. Before release, existing "
    "credentials must be invalidated or migrated through a forced reset and stored using Argon2id "
    "(preferred) or bcrypt with a unique salt; plaintext values and test bypasses must be removed."
))
replace_paragraph_starting(document, "The normal prototype performs analysis locally", (
    "The packaged clinical build uses a local SQLite database and binds its Flask service to the loopback "
    "interface, reducing external network exposure. However, exported reports are plaintext HTML files "
    "written to the Desktop or temporary directory, and the report template references an external font "
    "service. Protected reports must be encrypted or placed in an encrypted container, external resources "
    "must be removed from clinical exports, and any future synchronization must use TLS with institutionally "
    "managed credentials and explicit approval."
))
replace_paragraph_after_heading(document, "4.5 Authentication", (
    "As built, PsyClick provides clinician ID/password login and a password confirmation dialog for selected "
    "destructive UI actions. It does not establish a server-side authenticated session, does not issue or "
    "rotate a session token, and does not implement a minimum password policy, secure hashing, rate limiting, "
    "lockout, inactivity timeout, or MFA. A development verification bypass is also present in the reviewed "
    "API source and must be removed. The authentication policy in this document is therefore a required "
    "hardening baseline rather than an implemented control."
))
replace_paragraph_starting(document, "Authorization uses Role-Based Access Control", (
    "Target authorization uses Role-Based Access Control combined with record-level ownership rules. The "
    "current build does not implement backend roles or permission middleware: the frontend route guard only "
    "checks in-memory user state, while API endpoints accept caller-supplied clinician or patient identifiers. "
    "All protected operations—including record viewing, export, deletion, audit access, normative functions, "
    "and administration—must enforce the authenticated identity and deny access by default on the server side."
))
replace_paragraph_starting(document, "The penetration test is limited", (
    "No live penetration test was executed for this document. The following cases are an authorized test plan "
    "for an isolated laboratory copy containing synthetic data. Results must not be marked Passed until "
    "evidence such as request/response captures, logs, screenshots, and remediation retest records is attached."
))
replace_paragraph_starting(document, "The strongest existing security properties", (
    "The strongest evidenced properties are loopback-only service binding, Electron context isolation with "
    "Node integration disabled, local SQLite processing, pseudonymous client identifiers, parameterized database "
    "queries, transient analytical state, non-storage of full response text, an audit-event table, and a "
    "clinician-facing confidential report label. The thesis reports a Security evaluation mean of 4.64/5.00, "
    "but that perception-based result does not substitute for technical verification of password hashing, "
    "authorization, encryption, audit integrity, or recovery."
))
replace_paragraph_starting(document, "The primary residual risks", (
    "Static review identified Critical exposure from plaintext password storage and absent backend authentication/"
    "authorization. High risks include consent bypass through direct API use, unencrypted clinical files, mutable "
    "and unauthenticated audit functions, unprotected destructive endpoints, broad CORS, configuration-secret "
    "handling, and lack of verified backup/restore. These findings require remediation and independent retesting "
    "before institutional pilot deployment. Clinical automation bias remains a separate residual safety risk."
))
replace_paragraph_starting(document, "PsyClick demonstrates that a local", (
    "PsyClick demonstrates a functional and explainable clinical decision-support pipeline integrating hardware "
    "normalization, sliding-window filtering, within-session EWMA, Hotelling’s T², PSI/PAI contribution analysis, "
    "fuzzy classification, dashboards, reports, and audit events. The thesis documents 105 normative participants, "
    "15 clinical sessions, an overall ISO/IEC 25010 mean of 4.54/5.00, and Cohen’s κ = 0.466 (Moderate Agreement). "
    "These results support technical and research feasibility, not production security certification."
))
replace_paragraph_starting(document, "The system’s security posture benefits", (
    "The present build is suitable only for controlled demonstration or research use with synthetic or expressly "
    "approved data and strong workstation safeguards. It is not ready for unsupervised institutional or clinical "
    "production deployment because mandatory authentication, authorization, encryption, integrity, monitoring, "
    "and recovery controls remain incomplete. Implementing and independently verifying the Critical and High "
    "remediations is a release gate, while the output must remain non-diagnostic and subject to professional judgment."
))
replace_paragraph_starting(document, "PsyClick is a layered desktop application", (
    "PsyClick’s implemented application consists of an Electron/React interface, a loopback Flask API, telemetry "
    "capture and hardware-abstraction components, the feature-extraction and statistical analysis pipeline, an "
    "HTML report generator, an audit-event function, and a local SQLite datastore. Authentication, authorization, "
    "cryptographic, tamper-evident audit, monitoring, and backup services shown in the target architecture are "
    "required security components and are not all present in the current build. The analytical flow remains: "
    "normalized telemetry → sliding-window features → within-session EWMA baseline → Hotelling’s T² → Feature "
    "Contribution Analysis → PSI/PAI → fuzzy classification → Green/Amber/Red decision-support flag."
))
replace_paragraph_starting(document, "The PsyClick prototype demonstrates a strong foundational", (
    "The PsyClick prototype is functionally mature but security-incomplete. Thesis evaluators reported an "
    "ISO/IEC 25010 Security mean of 4.64/5.00 and overall quality mean of 4.54/5.00; static source review confirms "
    "several useful foundations, including local processing, loopback service binding, parameterized queries, "
    "pseudonymous identifiers, transient feature processing, audit-event capture, and Electron renderer isolation. "
    "The same review found Critical password and authorization weaknesses plus High-priority gaps in consent, "
    "encryption, report protection, audit integrity, secret management, and recovery."
))
replace_paragraph_starting(document, "Current thesis prototype:", (
    "Overall rating as of 10 July 2026: NOT READY FOR PRODUCTION CLINICAL DEPLOYMENT. The system may be used for "
    "controlled academic demonstration or approved research with synthetic/de-identified data and compensating "
    "endpoint controls. Release to an institutional pilot requires closure and retest of all Critical and High "
    "findings, privacy and clinical-owner approval, a successful restore exercise, and an authorized penetration test."
))

# Add the API documentation and make the operational guide an explicit user manual.
architecture_body = next(p for p in document.paragraphs if p.text.startswith("PsyClick’s implemented application consists"))
api_heading = add_paragraph_after(document, architecture_body, "API Documentation (Current Prototype)", "Heading 2")
api_note = add_paragraph_after(document, api_heading, "The desktop client calls a loopback Flask API under the /api prefix. The routes below describe the supplied prototype surface. They must be placed behind authenticated server middleware before any institutional deployment.")
api_rows = [
    ("POST /api/login", "Clinician credentials", "Login result; audit event"),
    ("POST /api/register", "Name and password", "Clinician account result"),
    ("POST /api/verify-clinician", "Clinician credentials", "Re-authentication result"),
    ("POST /api/logout", "None", "Logout audit event"),
    ("GET /api/stats", "Optional clinician_id", "Dashboard counts"),
    ("GET /api/sessions/recent", "Optional clinician_id", "Recent sessions"),
    ("GET /api/patients", "Optional clinician_id", "Pseudonymous clients"),
    ("GET /api/patients/{id}/sessions", "Client identifier", "Client session history"),
    ("POST /api/intake/start", "Patient and clinician identifiers", "Session context"),
    ("POST /api/calibration/*", "Session state", "Keyboard/mouse calibration"),
    ("POST /api/assessment/*", "Assessment payloads", "Feature capture and audit events"),
    ("GET /api/session/{id}", "Session identifier", "Session report data"),
    ("GET /api/audit", "Optional actor filter", "Audit events"),
    ("POST /api/export/report", "Report payload", "HTML report path"),
    ("DELETE /api/clients/{id}", "Client identifier", "Deleted session count"),
    ("GET /api/ping and /api/db-health", "None", "Readiness and database status"),
]
add_table_after(document, api_note, ["Endpoint", "Inputs", "Output / Purpose"], api_rows, widths=[2.0, 2.0, 3.0])
operational_heading = next(p for p in document.paragraphs if p.text.strip() == "Operational Guide")
operational_heading.text = "USER MANUAL – OPERATIONAL GUIDE"

# Recast target-only tables and add evidence-based statuses.
find_table(document, "Module", "Major Features").cell(0, 1).text = "Required / Target Features"
find_table(document, "Element", "Policy").cell(0, 1).text = "Proposed Policy"
find_table(document, "Permission", "Admin").cell(0, 0).text = "Target Permission"
legend_table = find_table(document, "Label", "Meaning")
legend_table.cell(2, 0).text = "Evidenced Current Control"
legend_table.cell(2, 1).text = (
    "Control directly observed in the reviewed application, such as loopback binding, parameterized queries, "
    "local SQLite processing, audit-event capture, or Electron renderer isolation; scope and limitations remain explicit."
)

control_rows = [
    ("Consent Gate", "Partial", "UI checkbox disables Start; backend start route does not persist or revalidate consent.", "Create versioned consent records and enforce a fail-closed backend gate before every capture start."),
    ("Least Privilege / RBAC", "Not Implemented", "Frontend has a basic logged-in guard; API has no authenticated roles, permissions, or ownership checks.", "Add server authentication middleware, RBAC, record ownership, and deny-by-default tests."),
    ("Secure Password Storage", "Not Implemented – Critical", "Clinician passwords are stored and compared as plaintext.", "Remove plaintext and test bypasses; force reset; use Argon2id/bcrypt and generic errors."),
    ("Input and SQL Safety", "Partially Implemented", "Many queries use parameters and IDs receive basic parsing; centralized request schemas are absent.", "Add type, length, range, state, Unicode, and file/path validation with negative tests."),
    ("Session Security", "Not Implemented – Critical", "No server session token, expiry, inactivity timeout, logout invalidation, or re-auth policy.", "Issue cryptographically random server sessions; bind identity/role; rotate, expire, and revoke tokens."),
    ("Data Encryption", "Not Implemented", "SQLite and exported HTML reports are plaintext at the application layer.", "Use encrypted database/fields, protected key storage, encrypted exports, and rotation/recovery procedures."),
    ("Audit Logging", "Partial", "Audit table and UI exist; entries are mutable and endpoints lack authorization and integrity chaining.", "Authenticate events; record user/object/outcome; chain hashes; restrict writes; alert and review."),
    ("Integrity Verification", "Not Implemented", "No report or backup hash/signature verification was found.", "Generate and verify SHA-256/HMAC or signatures for finalized reports and encrypted backups."),
    ("Privacy Minimization", "Partially Implemented", "Pseudonymous client IDs and local processing are used; full response text is not persisted, but raw key values exist transiently.", "Zero transient buffers promptly; document retention; implement withdrawal, deletion, and privacy tests."),
    ("Backup and Recovery", "Not Implemented", "No scheduled encrypted backup, inventory, hash verification, or restore function is evidenced.", "Implement 3-2-1-aligned encrypted backups, RPO/RTO monitoring, and quarterly restore tests."),
    ("Endpoint / Electron", "Partially Implemented", "API is loopback-only with debug off; Node integration is off and context isolation on; CORS and openExternal are broad.", "Restrict CORS/origin and IPC URL schemes; add CSP; patch/sign builds; keep full-disk encryption."),
]
replace_table_rows(
    find_table(document, "Control", "Implementation"),
    ["Control", "As-Built Status", "Reviewed Evidence", "Required Closure"],
    control_rows,
)
controls_table = find_table(document, "Control", "As-Built Status")
for row in controls_table.rows[1:]:
    add_status_color(row.cells[1], row.cells[1].text)

# Give functional test entries honest execution status.
functional_status = [
    "Partially evidenced by source review; successful login is thesis/prototype evidence, but secure-session controls fail.",
    "Partial: invalid credentials are denied, but generic errors, throttling, lockout, and enumeration resistance are incomplete.",
    "Failed by source review: UI gate exists, but the backend does not independently verify persisted consent.",
    "Validated functionally by thesis and implemented pipeline; security of result persistence remains separate.",
    "Partially evidenced: analytical objects reset, but memory zeroization and retention tests were not performed.",
    "Failed by source review: no backend record-ownership authorization is enforced.",
    "Partial: HTML export and confidentiality label exist; encryption, hash, and protected destination are absent.",
    "Not implemented / not tested.",
    "Not implemented / not tested.",
    "Not tested; SQLite transaction use provides only partial design evidence.",
]
ft = find_table(document, "Test ID", "Scenario")
ft.add_column(Inches(1.8))
set_cell_text(ft.rows[0].cells[3], "Observed Status", bold=True, color=WHITE)
for row, status in zip(ft.rows[1:], functional_status):
    set_cell_text(row.cells[3], status)
format_table(ft)
for row in ft.rows[1:]:
    add_status_color(row.cells[3], row.cells[3].text)

risk_rows = [
    ("Plaintext clinician password disclosure", "4", "5", "20 Critical", "Argon2id/bcrypt migration, forced reset, remove bypasses, protect credential store."),
    ("Unauthenticated API / privilege escalation", "4", "5", "20 Critical", "Server sessions, RBAC, ownership checks, deny-by-default middleware and tests."),
    ("Unauthorized local access to clinical records", "4", "5", "20 Critical", "Backend authorization, full-disk/database encryption, OS lock, access logs."),
    ("Consent bypass through direct API request", "3", "5", "15 High", "Versioned persisted consent and fail-closed backend check before capture."),
    ("Loss or theft of workstation or media", "3", "5", "15 High", "Full-disk/database encryption, encrypted backups, screen lock, physical controls."),
    ("Plaintext report export or temporary file", "4", "4", "16 High", "Encrypted/protected export, safe destination, no external resources, export audit."),
    ("Configuration secret exposure", "3", "5", "15 High", "Revoke exposed credentials, purge from packages/history, use OS/institution secret storage."),
    ("Malware or ransomware", "3", "5", "15 High", "Least privilege, endpoint protection, patching, application allowlisting, offline backups."),
    ("Audit forgery, deletion, or disclosure", "3", "4", "12 High", "Authenticated structured events, append-only/hash chain, restricted access, review alerts."),
    ("Data corruption with no verified recovery", "3", "5", "15 High", "Transactions, daily encrypted backups, hashes, isolated quarterly restore test."),
    ("Clinical over-reliance / automation bias", "3", "4", "12 High", "Non-diagnostic warnings, independent clinician assessment, training and supervision."),
    ("Local denial of service / resource exhaustion", "2", "3", "6 Moderate", "Request and input limits, throttling, disk monitoring, graceful recovery."),
]
replace_table_rows(
    find_table(document, "Risk", "Likelihood"),
    ["Risk", "Likelihood", "Impact", "Score", "Treatment"],
    risk_rows,
)
risk_table = find_table(document, "Risk", "Likelihood")
for row in risk_table.rows[1:]:
    add_status_color(row.cells[3], row.cells[3].text)

finding_rows = [
    ("Loopback API, debug disabled, and isolated Electron renderer", "Strength", "Maintain", "Retain 127.0.0.1 binding, Node integration off, context isolation on; add CSP and IPC restrictions."),
    ("Parameterized SQLite access in reviewed data functions", "Strength", "Maintain", "Keep prepared statements and add automated injection regression tests."),
    ("Clinician passwords stored and compared as plaintext", "Critical Gap", "Critical", "Invalidate/migrate credentials, remove test bypasses, and deploy Argon2id/bcrypt immediately."),
    ("No authenticated server session or backend RBAC", "Critical Gap", "Critical", "Implement session middleware, roles, permissions, ownership, expiry, revocation, and negative tests."),
    ("Consent gate enforced only by the UI", "Gap", "High", "Persist consent version/decision and require backend validation before any logger starts."),
    ("Database and clinical exports are not application-encrypted", "Gap", "High", "Encrypt restricted data and exports; protect keys outside the database."),
    ("Audit log is mutable and audit endpoints are unprotected", "Gap", "High", "Authenticate, restrict, hash-chain, monitor, and periodically review audit events."),
    ("Backup/restore implementation and evidence are absent", "Gap", "High", "Create encrypted backups, verify hashes, document RPO/RTO, and complete a restore drill."),
    ("Configuration contains a plaintext remote-service credential", "Critical Hygiene Gap", "High", "Revoke/rotate it, remove it from artifacts/history, and use protected secret injection."),
    ("Unrestricted CORS on a localhost API", "Gap", "High", "Allow only the Electron origin or remove browser CORS exposure; add anti-CSRF/origin checks as applicable."),
    ("Report export references an external font service", "Privacy Gap", "Moderate", "Bundle fonts locally and prohibit network requests from confidential reports."),
    ("Thesis security mean of 4.64/5.00", "Evaluation Evidence", "Context", "Report as expert perception/usability evidence; do not treat as penetration-test certification."),
]
replace_table_rows(
    find_table(document, "Finding", "Status"),
    ["Finding", "Classification", "Priority", "Required Action"],
    finding_rows,
)
findings_table = find_table(document, "Finding", "Classification")
for row in findings_table.rows[1:]:
    add_status_color(row.cells[2], row.cells[2].text)

# Add a final evidence appendix.
last = document.paragraphs[-1]
appendix = add_paragraph_after(document, last, "APPENDIX B – SOURCE-CODE SECURITY REVIEW", "Heading 1")
appendix.paragraph_format.page_break_before = True
review_scope = add_paragraph_after(
    document,
    appendix,
    "Review date: 10 July 2026. Method: static review of the supplied Electron/React frontend, packaged Python "
    "modules, database schema and access functions, report exporter, configuration, and the corresponding tracked "
    "API source. This review did not execute a live attack, inspect a production workstation, validate operating-"
    "system permissions, or process real patient data. Findings are therefore code-evidence conclusions and must "
    "be followed by controlled dynamic testing after remediation.",
)
evidence_rows = [
    ("Local service exposure", "API startup configuration", "Binds to 127.0.0.1:5001; debug disabled.", "Implemented foundation"),
    ("Electron renderer isolation", "electron/main.js", "nodeIntegration false; contextIsolation true.", "Implemented foundation"),
    ("Credential storage", "database_manager.py clinician schema and verification", "Password field is plaintext and compared directly.", "Critical failure"),
    ("Session management", "frontend context/API client and API routes", "No server-issued session token or authenticated request context.", "Critical failure"),
    ("Authorization", "API record, audit, export, and delete routes", "No role/permission middleware; identifiers are caller supplied.", "Critical failure"),
    ("Consent enforcement", "Intake UI and intake/start route", "Checkbox controls the UI only; backend receives no consent artifact.", "High gap"),
    ("SQL injection resistance", "database manager and API query calls", "Values are generally passed as parameters; fixed fragments build optional filters.", "Partially implemented"),
    ("Data at rest", "SQLite configuration and export routines", "No database/field encryption; HTML reports are plaintext.", "High gap"),
    ("Auditability", "audit_log schema, log/read endpoints and Audit page", "Events are recorded and viewable; no authentication, chaining, or immutability.", "Partial / High gap"),
    ("Backup and recovery", "Application modules and API surface", "No evidenced scheduled backup, integrity verification, or restore workflow.", "High gap"),
    ("Telemetry minimization", "backend controller snapshots", "Full response text is not persisted; response length/features are stored; raw key values are transient.", "Partially implemented"),
    ("Report privacy", "report_exporter.py", "Confidential label present; external font reference and unencrypted Desktop/temp output remain.", "Partial / Moderate gap"),
]
evidence_table = add_table_after(
    document,
    review_scope,
    ["Review Area", "Evidence Location", "Observation", "Conclusion"],
    evidence_rows,
    widths=[1.25, 1.9, 3.2, 1.15],
)
for row in evidence_table.rows[1:]:
    add_status_color(row.cells[3], row.cells[3].text)

gate_heading = add_paragraph_after(document, evidence_table.rows[-1].cells[-1].paragraphs[-1], "RELEASE GATE", "Heading 2")
# The previous insertion is inside a cell; move the paragraph after the table.
evidence_table._tbl.addnext(gate_heading._p)
gate_text = add_paragraph_after(
    document,
    gate_heading,
    "PsyClick shall not be represented as secure for production clinical deployment until: (1) plaintext "
    "credentials and bypasses are removed; (2) authenticated server sessions, RBAC, and ownership checks pass "
    "negative tests; (3) backend consent enforcement is verified; (4) restricted data, reports, and backups are "
    "encrypted with protected keys; (5) audit records are tamper-evident and access-controlled; (6) a backup is "
    "successfully restored and verified; (7) exposed secrets are rotated and removed; and (8) an authorized "
    "independent vulnerability assessment and penetration test reports no unresolved Critical or High finding.",
)

# Trace every required final-project deliverable and grading area to evidence and a submission status.
matrix_heading = add_paragraph_after(document, gate_text, "APPENDIX C – FINAL PROJECT DELIVERABLE COMPLIANCE MATRIX", "Heading 1")
matrix_heading.paragraph_format.page_break_before = True
matrix_intro = add_paragraph_after(document, matrix_heading, "This matrix answers the CS0029 specification directly. “Complete in this document” means the section or evidence is present here; “Separate artifact” identifies a submission item that cannot be created by a security document alone, such as the executable application, source repository, or defense presentation.")
deliverable_rows = [
    ("Complete thesis manuscript, Chapters 1–6", "Chapters 1–6 coverage now appears in this document; full thesis manuscript remains the authoritative research submission.", "Partial / separate thesis required", "Submit the final thesis PDF/DOCX."),
    ("Working system prototype", "PsyClick Electron/React frontend and packaged Python API are present in the supplied workspace.", "Separate artifact – available", "Submit source repository and tested build/package."),
    ("User Management module", "Clinician registration/login exists; role lifecycle and secure password controls are incomplete.", "Partial", "Implement hashing, lifecycle, roles, lockout, MFA."),
    ("Authentication module", "Login and re-authentication UI/API exist; server sessions and secure credential handling are incomplete.", "Partial / Critical gap", "Close release-gate authentication findings."),
    ("Dashboard module", "Dashboard, session history, patient/client views, and report navigation are implemented.", "Complete functionally", "Demonstrate with synthetic data."),
    ("Main business-process module", "Intake, consent UI, keyboard/mouse calibration, PHQ-9, GAD-7, emotional task, and analytics are implemented.", "Complete functionally / security hardening required", "Demonstrate consent and assessment workflow."),
    ("Reports module", "HTML report and summary export exist with clinical rationale and confidentiality label.", "Partial", "Encrypt exports, remove external resources, hash and log exports."),
    ("Audit Logs module", "Audit table, API, and Audit page exist.", "Partial", "Authenticate, restrict, chain, monitor, and review logs."),
    ("Threat model", "Assets, actors, STRIDE summary, and threat controls are documented.", "Complete in this document", "Use in defense and security report."),
    ("Risk assessment matrix", "5×5 matrix plus prioritized as-built risks and treatments are documented.", "Complete in this document", "Obtain owner acceptance after remediation."),
    ("Security policy", "Purpose, scope, access, privacy, encryption, logging, backup, and non-diagnostic policy statements are documented.", "Complete in this document", "Approve institutionally before deployment."),
    ("Incident response plan", "Roles, phases, severity guide, containment, recovery, and post-incident review are documented.", "Complete in this document", "Run tabletop exercise."),
    ("Backup and recovery plan", "Backup strategy, RPO/RTO, restore procedure, and recovery controls are documented.", "Plan complete; implementation unverified", "Implement encrypted backups and complete restore drill."),
    ("Architecture diagram", "Layered security architecture and local-first network diagram are embedded.", "Complete in this document", "Use figures in thesis/presentation."),
    ("DFD and ERD", "Context/Level-1 DFD and core ERD are embedded; database tables are documented.", "Complete in this document", "Validate against final schema."),
    ("API documentation", "Prototype endpoint table is included under Technical Documentation.", "Complete in this document", "Update after security middleware is implemented."),
    ("Installation guide", "Proposed secure deployment installation steps are documented.", "Complete as proposed guide", "Validate on a clean workstation."),
    ("User manual", "User Manual – Operational Guide covers clinician workflow, consent, assessment, reports, and logout.", "Complete in this document", "Add screenshots if required by instructor."),
    ("Security evaluation report", "Executive summary, findings, rating, functional/security test plans, and source review are documented.", "Complete with limitations", "Attach dynamic test evidence after remediation."),
    ("Presentation slides", "A 20-slide presentation outline is documented.", "Outline only / separate artifact", "Create and submit slide deck."),
    ("Source code repository", "The supplied workspace is the implementation source.", "Separate artifact – available", "Submit repository URL/archive and commit history."),
    ("Final defense demonstration", "Demonstration sequence is described in the presentation outline and user manual.", "Separate activity", "Perform live login, consent, workflow, report, audit, and recovery demo."),
]
deliverable_table = add_table_after(document, matrix_intro, ["Specification Deliverable", "Evidence", "Status", "Action / Submission"], deliverable_rows, widths=[1.8, 3.2, 1.25, 1.25])
for row in deliverable_table.rows[1:]:
    add_status_color(row.cells[2], row.cells[2].text)

rubric_heading = add_paragraph_after(document, deliverable_table.rows[-1].cells[-1].paragraphs[-1], "Rubric Evidence Summary (100 Points)", "Heading 2")
deliverable_table._tbl.addnext(rubric_heading._p)
rubric_rows = [
    ("Documentation – 25", "Organization and structure; literature review; technical content; security documentation; writing quality.", "Chapters 1–6 coverage, literature synthesis, diagrams, controls, policies, risk matrix, and appendices.", "Strong coverage; full thesis and references must be submitted separately."),
    ("System Design & Development – 25", "Functionality; UI design; database design; technical innovation.", "Implemented dashboard/workflow, SQLite schema, analytics pipeline, local-first architecture, and prototype modules.", "Functionally evidenced; verify clean-build behavior and database permissions."),
    ("Information Assurance & Security – 30", "Authentication; authorization; data protection; risk; testing; audit/monitoring.", "Threat model, matrix, policies, test plans, audit module, and static findings.", "Not fully achieved until Critical/High gaps are remediated and dynamically retested."),
    ("Final Defense & Presentation – 20", "Presentation quality; demonstration; completeness; professionalism/teamwork.", "20-slide outline, demo sequence, findings, recommendations, and release gate.", "Requires separate slide deck, rehearsal, live demonstration, and team defense."),
]
rubric_table = add_table_after(document, rubric_heading, ["Rubric Area", "Criteria", "Evidence", "Completion Position"], rubric_rows, widths=[1.3, 2.2, 2.5, 1.5])
for row in rubric_table.rows[1:]:
    add_status_color(row.cells[3], row.cells[3].text)

# Improve pagination and headings without rewriting the source layout.
break_titles = {
    "CHAPTER 1 – INTRODUCTION",
    "CHAPTER 2 – REVIEW OF RELATED LITERATURE AND STUDIES",
    "CHAPTER 3 – SYSTEM ANALYSIS AND DESIGN",
    "CHAPTER 4 – SYSTEM DEVELOPMENT AND SECURITY IMPLEMENTATION",
    "CHAPTER 5 – SECURITY EVALUATION AND TESTING",
    "CHAPTER 6 – CONCLUSIONS AND RECOMMENDATIONS",
    "THREAT MODEL",
    "SECURITY POLICY",
    "INCIDENT RESPONSE PLAN",
    "BACKUP AND RECOVERY PLAN",
    "TECHNICAL DOCUMENTATION",
    "SECURITY EVALUATION REPORT",
    "PRESENTATION OUTLINE",
    "APPENDIX A – IMPLEMENTATION STATUS LEGEND",
    "APPENDIX B – SOURCE-CODE SECURITY REVIEW",
    "APPENDIX C – FINAL PROJECT DELIVERABLE COMPLIANCE MATRIX",
}
for p in document.paragraphs:
    if p.text.strip() in break_titles:
        p.paragraph_format.page_break_before = True
    if p.style.name.startswith("Heading"):
        p.paragraph_format.keep_with_next = True

# Standardize major tables while preserving content.
for table in document.tables:
    format_table(table)

# Footer with document label and live page-number field.
section = document.sections[0]
footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.text = "PsyClick – CS0029 Information Assurance and Security  |  Page "
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer_p.add_run()
fld_char_begin = OxmlElement("w:fldChar")
fld_char_begin.set(qn("w:fldCharType"), "begin")
instr_text = OxmlElement("w:instrText")
instr_text.set(qn("xml:space"), "preserve")
instr_text.text = " PAGE "
fld_char_end = OxmlElement("w:fldChar")
fld_char_end.set(qn("w:fldCharType"), "end")
run._r.extend([fld_char_begin, instr_text, fld_char_end])
for footer_run in footer_p.runs:
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(90, 105, 120)

document.save(OUTPUT)
print(OUTPUT)
