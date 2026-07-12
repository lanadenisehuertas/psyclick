from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "artifacts" / "_work"
WORK.mkdir(parents=True, exist_ok=True)


def extract_docx(source: Path, output: Path) -> None:
    document = Document(source)
    lines = [
        f"SOURCE: {source}",
        f"PARAGRAPHS: {len(document.paragraphs)}",
        f"TABLES: {len(document.tables)}",
        "",
        "=== PARAGRAPHS ===",
    ]
    for index, paragraph in enumerate(document.paragraphs):
        lines.append(f"P{index} [{paragraph.style.name}]: {paragraph.text}")
    lines.append("")
    lines.append("=== TABLES ===")
    for table_index, table in enumerate(document.tables):
        lines.append(
            f"TABLE {table_index} ({len(table.rows)} rows x {len(table.columns)} columns)"
        )
        for row_index, row in enumerate(table.rows):
            cells = [cell.text.replace("\n", " / ") for cell in row.cells]
            lines.append(f"R{row_index}: " + " | ".join(cells))
        lines.append("")
    output.write_text("\n".join(lines), encoding="utf-8")


def extract_pdf(source: Path, output: Path) -> None:
    reader = PdfReader(source)
    lines = [f"SOURCE: {source}", f"PAGES: {len(reader.pages)}", ""]
    for page_index, page in enumerate(reader.pages, start=1):
        lines.append(f"=== PAGE {page_index} ===")
        lines.append(page.extract_text() or "")
        lines.append("")
    output.write_text("\n".join(lines), encoding="utf-8")


extract_docx(
    Path(r"C:\Users\Lana\Downloads\Final Project Specification.docx"),
    WORK / "final_project_specification.txt",
)
extract_docx(
    Path(r"C:\Users\Lana\Downloads\PsyClick_CS0029_Security_Document.docx"),
    WORK / "psyclick_security_document.txt",
)
extract_pdf(
    Path(r"C:\Users\Lana\Downloads\BYTEME-Thesis Paper.pdf"),
    WORK / "byteme_thesis_paper.txt",
)
