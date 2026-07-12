from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptInches, Pt as PptPt


ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "artifacts"
DIAGRAMS = ART / "diagrams"
ART.mkdir(parents=True, exist_ok=True)


def create_user_manual():
    out = ART / "PsyClick_CS0029_User_Manual.docx"
    doc = Document()
    doc.core_properties.title = "PsyClick CS0029 User Manual"
    doc.core_properties.subject = "Clinician and administrator operating guide"
    doc.add_heading("PsyClick Clinical Decision Support System", 0)
    p = doc.add_paragraph("CS0029 Information Assurance and Security – User Manual")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Version: July 2026 | For controlled academic demonstration and approved research use")
    doc.add_heading("Important Safety and Privacy Notice", 1)
    doc.add_paragraph("PsyClick is non-diagnostic decision support. It does not replace clinical judgment, risk assessment, or emergency procedures. Use only approved synthetic, de-identified, or consented data. Do not treat Green/Amber/Red flags as diagnoses.")
    doc.add_heading("1. Start the Application", 1)
    for item in [
        "Use a supported workstation with current patches, endpoint protection, automatic screen lock, and restricted physical access.",
        "Launch the packaged PsyClick application. Confirm the API readiness indicator before signing in.",
        "Use an individual clinician account. Never share credentials or use the development/test bypass.",
    ]: doc.add_paragraph(item, style="List Number")
    doc.add_heading("2. Sign In and Verify Identity", 1)
    doc.add_paragraph("Enter the clinician ID and password. A successful login opens the dashboard and records an audit event. Invalid attempts must be reported if repeated. In the current prototype, server-side sessions, lockout, MFA, and secure password hashing remain hardening requirements; do not deploy with real clinical data until those controls are verified.")
    doc.add_heading("3. Create a Client Context and Record Consent", 1)
    for item in [
        "Create or select a pseudonymous client identifier; avoid unnecessary names, addresses, or other direct identifiers.",
        "Explain the purpose, features captured, retention, limits, and non-diagnostic nature of the system.",
        "Record the participant’s informed consent before starting telemetry. The supplied UI has a consent checkbox; the backend consent record and fail-closed validation are required before production use.",
    ]: doc.add_paragraph(item, style="List Number")
    doc.add_heading("4. Run Calibration and Assessment", 1)
    for item in [
        "Run keyboard calibration, then mouse calibration, using the guided prompts.",
        "Administer PHQ-9 and GAD-7 according to the approved study/clinical protocol.",
        "Run the emotional-task prompts. Supervise the participant and stop if distress or technical instability occurs.",
        "Allow the system to finalize the session. Do not close the application during analysis or database writes.",
    ]: doc.add_paragraph(item, style="List Number")
    doc.add_heading("5. Interpret the Dashboard and Report", 1)
    doc.add_paragraph("Review T², threshold, PSI, PAI, confidence, domain/level profiles, and the plain-language rationale. Compare the output with independent clinical observation. Record contextual factors and never infer a diagnosis from a flag alone.")
    doc.add_heading("6. Export and Handle Reports", 1)
    for item in [
        "Export only when necessary and only to an approved protected destination.",
        "Treat reports as restricted clinical/behavioral information. Do not email or upload unencrypted reports.",
        "Verify that the export is classified and logged. The current prototype’s plaintext HTML export must be encrypted before institutional deployment.",
    ]: doc.add_paragraph(item, style="List Number")
    doc.add_heading("7. Audit, Logout, and Close", 1)
    for item in [
        "Use the Audit page to review login, consent, access, export, configuration, and backup events.",
        "Log out after the session. Lock the workstation whenever it is unattended.",
        "Discard or securely delete temporary files and follow the approved retention schedule.",
    ]: doc.add_paragraph(item, style="List Number")
    doc.add_heading("8. Administrator and Recovery Tasks", 1)
    doc.add_paragraph("Administrators must provision least-privilege accounts, review failed logins and unusual exports, verify encrypted backups, perform restore drills, and activate the incident-response plan for suspected disclosure, malware, consent failure, corruption, or unauthorized access. The current prototype does not yet evidence a complete backup/restore implementation; this is a release gate.")
    doc.add_heading("Troubleshooting", 1)
    doc.add_paragraph("If the API is unavailable, restart the application and confirm the packaged backend is running on the loopback interface. If a session fails, preserve the audit information, do not reuse an uncertain partial report, and notify the system administrator. Do not expose database files or credentials while troubleshooting.")
    doc.add_heading("Support and Escalation", 1)
    doc.add_paragraph("Report security incidents immediately to the Incident Coordinator, System Administrator, Data Protection Officer/Privacy Lead, and Clinical Lead using the response plan in the CS0029 Security Document.")
    for section in doc.sections:
        section.footer.paragraphs[0].text = "PsyClick – CS0029 User Manual"
    doc.save(out)
    return out


def add_slide(prs, title, bullets=None, image=None, footer="PsyClick | CS0029 | July 2026"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = PptRGBColor(245, 248, 250)
    title_box = slide.shapes.add_textbox(PptInches(0.55), PptInches(0.35), PptInches(12.2), PptInches(0.65))
    tf = title_box.text_frame; tf.clear(); p = tf.paragraphs[0]; p.text = title; p.font.size = PptPt(25); p.font.bold = True; p.font.color.rgb = PptRGBColor(23,50,77)
    if image:
        slide.shapes.add_picture(str(image), PptInches(0.7), PptInches(1.2), width=PptInches(12.0), height=PptInches(5.7))
    if bullets:
        box = slide.shapes.add_textbox(PptInches(0.8), PptInches(1.35), PptInches(11.7), PptInches(5.4))
        tf = box.text_frame; tf.word_wrap = True; tf.clear()
        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = bullet; p.level = 0; p.font.size = PptPt(20 if len(bullets) < 5 else 16); p.font.color.rgb = PptRGBColor(40, 56, 70); p.space_after = PptPt(12)
    foot = slide.shapes.add_textbox(PptInches(0.6), PptInches(7.05), PptInches(12.0), PptInches(0.25)); p = foot.text_frame.paragraphs[0]; p.text = footer; p.font.size = PptPt(9); p.font.color.rgb = PptRGBColor(100, 120, 135); p.alignment = PP_ALIGN.RIGHT
    return slide


def create_presentation():
    out = ART / "PsyClick_CS0029_Final_Defense_Presentation.pptx"
    prs = Presentation(); prs.slide_width = PptInches(13.333); prs.slide_height = PptInches(7.5)
    add_slide(prs, "PsyClick", ["Clinical decision support for single-session psychomotor disturbances", "EWMA-based multivariate analysis of keystroke and cursor dynamics", "CS0029 Information Assurance and Security | July 2026"])
    add_slide(prs, "Background and Problem", ["Mental-health assessment can miss distress because of concealment, recall bias, stigma, and delayed help-seeking.", "PsyClick supplies objective behavioral evidence while keeping professional judgment central.", "Security is essential because telemetry may function as behavioral biometric data."])
    add_slide(prs, "Proposed Solution", ["Local-first, non-diagnostic desktop workflow.", "10–15 minute session: login → intake/consent → calibration → tasks → analytics → dashboard/report.", "Data minimization, pseudonymous identifiers, and clinician-facing rationale."])
    add_slide(prs, "Objectives", ["Hardware abstraction and high-resolution telemetry.", "Within-session EWMA personal baseline.", "Hotelling’s T², PSI/PAI contribution analysis, fuzzy classification.", "Information-assurance controls and security evaluation."])
    add_slide(prs, "System Workflow", ["1. Authenticate clinician", "2. Create/select pseudonymous client and record consent", "3. Calibrate keyboard and mouse", "4. Administer PHQ-9/GAD-7 and emotional tasks", "5. Review report, export only when authorized, audit and logout"])
    add_slide(prs, "System Architecture", image=DIAGRAMS / "system-architecture.png")
    add_slide(prs, "Data Flow and ERD", image=DIAGRAMS / "dfd.png")
    add_slide(prs, "Database Model", image=DIAGRAMS / "erd.png")
    add_slide(prs, "Network Architecture", image=DIAGRAMS / "network-architecture.png")
    add_slide(prs, "Analytical Pipeline", ["HAL normalization → sliding-window features → EWMA baseline.", "Eight biomarkers: flight time, dwell time, typing velocity, error rate, path entropy, cursor velocity, jerk, pause frequency.", "Hotelling’s T² → Feature Contribution Analysis → PSI/PAI → fuzzy logic → Green/Amber/Red."])
    add_slide(prs, "Threat Model", ["Assets: identity, consent, telemetry, derived features, reports, credentials, keys, audit records, clinical trust.", "Actors: local insiders, compromised accounts, malware, accidental actors, supply-chain attackers.", "STRIDE: spoofing, tampering, repudiation, information disclosure, denial of service, elevation of privilege."])
    add_slide(prs, "Security Architecture", ["Target: authenticated sessions, RBAC, consent gate, encryption, integrity verification, audit chain, backup and recovery.", "Current prototype: loopback API, local SQLite, parameterized queries, Electron isolation, audit-event capture.", "High-priority gaps are explicitly tracked; security claims are not overstated."])
    add_slide(prs, "Authentication and Authorization", ["Current: clinician login and UI guard.", "Gaps: plaintext passwords, no server session token, no lockout/MFA, no backend RBAC or record ownership.", "Release gate: Argon2id/bcrypt, session middleware, roles, ownership checks, expiry, revocation, negative tests."])
    add_slide(prs, "Data Protection and Privacy", ["Current: local processing, pseudonymous client identifiers, non-storage of full response text.", "Gaps: plaintext SQLite/reports, transient raw key values, external report font reference, exposed configuration secret.", "Required: encryption, protected keys, retention/deletion, protected exports, secret rotation."])
    add_slide(prs, "Risk Assessment", ["Critical: plaintext clinician passwords; unauthenticated API/privilege escalation; unauthorized record access.", "High: consent bypass, plaintext exports, exposed configuration secret, audit tampering, no verified recovery.", "Risk acceptance: Critical before release; High requires treatment and owner approval."])
    add_slide(prs, "Security Testing", ["Functional cases: login, consent, calibration, analysis, export, backup, restore, interruption.", "Vulnerability review: dependencies, auth, authorization, database, paths, reports, logging, backups.", "Penetration tests are authorized laboratory simulations only; no live clinical records."])
    add_slide(prs, "Thesis Results", ["105 normative adult participants and 15 clinical sessions.", "Overall ISO/IEC 25010 mean: 4.54/5.00.", "Security evaluation mean: 4.64/5.00; Reliability: 4.67/5.00.", "Cohen’s κ = 0.466, Moderate Agreement."])
    add_slide(prs, "Incident Response and Recovery", ["Preparation → identification → containment → eradication → recovery → post-incident review.", "Proposed RPO: 24 hours; proposed RTO: 4 hours for a workstation deployment.", "Encrypted backups, hash verification, isolated restore, administrator and clinical-owner approval."])
    add_slide(prs, "Conclusions and Recommendations", ["PsyClick is a functional, explainable research prototype and behavioral signal amplifier.", "It is not production-clinical-ready until Critical/High security findings are closed and independently retested.", "Priorities: password/session/RBAC hardening, encryption, audit integrity, backup/restore, privacy procedures, testing, and longitudinal validation."])
    add_slide(prs, "Live Demonstration", ["Login and re-authentication", "Consent-gated intake and calibration", "Core assessment and dashboard", "Report export and audit review", "Backup verification / recovery procedure (synthetic data only)"])
    add_slide(prs, "Questions", ["Thank you.", "PsyClick remains non-diagnostic and subject to professional judgment."])
    prs.save(out)
    return out


def create_checklist():
    out = ART / "FINAL_PROJECT_DELIVERABLE_CHECKLIST.md"
    out.write_text("""# PsyClick CS0029 Final Project Submission Checklist\n\nGenerated 10 July 2026 from the Final Project Specification, thesis, and supplied system.\n\n## Prepared artifacts\n\n- [x] Completed security document: `PsyClick_CS0029_Security_Document_COMPLETED.docx`\n- [x] User manual: `PsyClick_CS0029_User_Manual.docx`\n- [x] Final defense deck: `PsyClick_CS0029_Final_Defense_Presentation.pptx`\n- [x] Architecture, DFD, ERD, and network figures: `diagrams/`\n- [x] Deliverable/rubric traceability matrix embedded in the security document\n- [x] Source-code review and release-gate findings embedded in the security document\n\n## Required separate submissions\n\n- [ ] Full thesis manuscript Chapters 1–6, with bibliography, appendices, and approved institutional formatting\n- [ ] Source code repository/archive and working prototype build\n- [ ] Demonstrated clean installation on the target workstation\n- [ ] Dynamic functional, vulnerability, and authorized penetration-test evidence\n- [ ] User acceptance test forms/signatures\n- [ ] Backup/restore drill evidence\n- [ ] Final defense rehearsal, live demonstration, and team presentation\n\n## Release blockers found in the supplied system\n\n- [ ] Replace plaintext clinician passwords with Argon2id/bcrypt and remove development bypasses\n- [ ] Add authenticated server sessions, RBAC, record ownership, expiry, revocation, and lockout/MFA\n- [ ] Enforce consent in the backend before any capture starts\n- [ ] Encrypt restricted database fields/files, reports, backups, and protect keys\n- [ ] Make audit events authenticated, append-only/hash-chained, and reviewable\n- [ ] Implement encrypted scheduled backups and a verified restore workflow\n- [ ] Remove/rotate exposed configuration secrets and restrict CORS/IPC/external report resources\n\n**Conclusion:** the documentation package is substantially complete, but the final project as a whole is not yet ready to claim full security completion until the separate artifacts and release blockers above are addressed and evidenced.\n""", encoding="utf-8")
    return out


if __name__ == "__main__":
    print(create_user_manual())
    print(create_presentation())
    print(create_checklist())
